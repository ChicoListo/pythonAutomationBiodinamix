from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

# Crear un nuevo libro de trabajo
wb = Workbook()

# Crear una hoja de cálculo para el reporte de calidad
sheet = wb.active
sheet.title = 'Reporte de Calidad'

# Ejemplo de datos de reporte (puedes modificar esta parte según tus necesidades)
productos = ['Producto A', 'Producto B', 'Producto C', 'Producto D']
cantidades_aceptables = [100, 150, 80, 120]
cantidades_defectuosas = [20, 30, 15, 25]

# Escribir datos en el archivo
sheet['A1'] = 'Producto'
sheet['B1'] = 'Cantidad Aceptable'
sheet['C1'] = 'Cantidad Defectuosa'

for idx, producto in enumerate(productos, start=2):
    sheet[f'A{idx}'] = producto
    sheet[f'B{idx}'] = cantidades_aceptables[idx - 2]
    sheet[f'C{idx}'] = cantidades_defectuosas[idx - 2]

# Crear gráfico de barras
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Cantidad Aceptable vs Defectuosa por Producto"
chart.x_axis.title = 'Productos'
chart.y_axis.title = 'Cantidad'

data = Reference(sheet, min_col=2, min_row=1, max_row=len(productos) + 1, max_col=3)
cats = Reference(sheet, min_col=1, min_row=2, max_row=len(productos) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

# Ubicar el gráfico dentro de la hoja de cálculo
sheet.add_chart(chart, "E1")

# Ajustar ancho de columnas automáticamente
for col in sheet.iter_cols():
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(cell.value)
        except:
            pass
    adjusted_width = (max_length + 2) * 1.2
    sheet.column_dimensions[column].width = adjusted_width

# Guardar el libro de trabajo
wb.save('reporte_calidad.xlsx')

# Cerrar el libro de trabajo
wb.close()

print("Archivo 'reporte_calidad.xlsx' creado correctamente.")
